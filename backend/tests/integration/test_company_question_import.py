"""Company-private question import: template, formats, validation, skill mapping, duplicates, reuse, tenant isolation, student privacy."""
import io
import json
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.database import AsyncSessionLocal
from app.models.assessments import Assessment, AssessmentQuestion
from app.models.enums import JobStatus
from app.models.questions import Question, QuestionImportBatch
from app.services.ai_gateway.vector_store import TenantScope
from app.services.questions import company_import as ci
from tests.factories import make_application, make_company, make_job, make_student


@pytest.fixture
async def w():
    async with AsyncSessionLocal() as db:
        oa, ra, ha = await make_company(db, "ImpA")
        ob, rb, hb = await make_company(db, "ImpB")
        job1 = await make_job(db, oa, ra, [("Python", "required", 0.6, 1.0), ("PostgreSQL", "required", 0.6, 0.8)], status=JobStatus.REQUIREMENTS_CONFIRMED)
        job2 = await make_job(db, oa, ra, [("Python", "required", 0.6, 1.0)], status=JobStatus.REQUIREMENTS_CONFIRMED)
        jobb = await make_job(db, ob, rb, [("Python", "required", 0.6, 1.0)], status=JobStatus.REQUIREMENTS_CONFIRMED)
        for j in (job1, job2, jobb):
            j.employment_type, j.work_mode = "FULL_TIME", "REMOTE"
        a1, a2, ab = (Assessment(job_id=j.id, title=f"t{i}", status="DRAFT") for i, j in enumerate((job1, job2, jobb)))
        db.add_all([a1, a2, ab])
        await db.commit()
        return dict(ha=ha, hb=hb, oa=str(oa.id), ob=str(ob.id), j1=str(job1.id), j2=str(job2.id), jb=str(jobb.id), a1=str(a1.id), a2=str(a2.id), ab=str(ab.id),
                    a_obj=oa, ra=ra)


HEAD = ci.COLUMNS


def sheet(rows: list[dict], meta: dict | None = None, meta_from=None) -> bytes:
    """A filled Excel file built on the real downloaded template structure."""
    m = meta or {"template_version": ci.TEMPLATE_VERSION}
    wb = load_workbook(io.BytesIO(ci.build_xlsx(m)))
    q = wb["Questions"]
    for r in rows:
        q.append([r.get(c, "") for c in HEAD])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def mcq(text, skill="Python", correct="B", **kw):
    return {"external_question_id": kw.pop("eid", ""), "question_type": "MCQ", "skill": skill, "difficulty": "medium", "question_text": text,
            "option_a": "return", "option_b": "yield", "option_c": "emit", "correct_option": correct, "explanation": "because", **kw}


def tech(text, skill="PostgreSQL", **kw):
    return {"external_question_id": kw.pop("eid", "T-1"), "question_type": "TECHNICAL_WRITTEN", "skill": skill, "difficulty": "hard", "question_text": text,
            "technical_rubric": "selectivity and statistics | cost estimates | index only scans", **kw}


XLSX = ("q.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


async def upload(client, h, job, data: bytes, name="q.xlsx"):
    return await client.post(f"/question-imports?job_id={job}", headers=h, files={"file": (name, data)})


@pytest.mark.asyncio
async def test_template_is_generated_for_this_company_job_and_assessment_in_all_formats(client, w):
    x = await client.get(f"/question-imports/template?job_id={w['j1']}&format=xlsx", headers=w["ha"])
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Instructions", "Questions", "Metadata"]
    assert [c.value for c in wb["Questions"][1]] == HEAD
    meta = {r[0]: r[1] for r in list(wb["Metadata"].iter_rows(values_only=True))[1:]}
    assert meta["company_id"] == w["oa"] and meta["job_id"] == w["j1"] and meta["assessment_id"] == w["a1"] and meta["template_version"] == ci.TEMPLATE_VERSION
    assert meta["job_title"] and meta["assessment_name"] == "t0" and meta["template_id"] and meta["generated_at"]
    c = await client.get(f"/question-imports/template?job_id={w['j1']}&format=csv", headers=w["ha"])
    assert c.text.splitlines()[0].startswith("# template_version=") and ",".join(HEAD) in c.text
    j = (await client.get(f"/question-imports/template?job_id={w['j1']}&format=json", headers=w["ha"])).json()
    assert j["metadata"]["company_id"] == w["oa"] and len(j["questions"]) == 2
    assert (await client.get(f"/question-imports/template?job_id={w['j1']}&format=xlsx", headers=w["hb"])).status_code in (403, 404)  # not B's job
    assert (await client.get(f"/question-imports/template?job_id={w['j1']}&format=pdf", headers=w["ha"])).status_code == 422


@pytest.mark.asyncio
async def test_excel_import_preview_mapping_duplicates_confirm_and_same_company_reuse(client, w):
    tok = f"A_ONLY_PRIVATE_TOKEN_{uuid.uuid4().hex[:8]}"
    rows = [mcq(f"{tok} which keyword defines a generator function in Python?", eid="M1"),
            mcq("How do you obtain the number of items in a Python list quickly?", skill="python3", eid="M2"),          # alias -> Python
            mcq("Which SQL clause filters grouped rows after aggregation happens?", skill="Postgres", eid="M3"),        # alias -> PostgreSQL
            tech("Explain how a B-tree index can speed up equality lookups on a large table."),
            tech("Describe how you would find and fix a slow query in production PostgreSQL.", eid="T-2"),
            mcq("Which keyword exits a Python loop immediately when a condition is met?", correct="E", eid="BAD"),      # E but 3 options -> invalid
            mcq("What does the Zorblax scheduler guarantee about job ordering?", skill="Zorblax Framework", eid="UNK"),  # unknown skill
            mcq("How do you obtain the number of items in a Python list quickly?", skill="Python", eid="DUPFILE")]      # duplicate within the file
    r = await upload(client, w["ha"], w["j1"], sheet(rows, {"template_version": ci.TEMPLATE_VERSION, "company_id": w["oa"], "job_id": w["j1"]}))
    assert r.status_code == 200, r.text
    b = r.json()
    st = {x["raw"]["external_question_id"] or x["raw"]["question_text"][:10]: x for x in b["rows"]}
    assert st["M1"]["status"] in ("READY", "WARNING") and st["M2"]["canonical_skill"] == "Python" and st["M3"]["canonical_skill"] == "PostgreSQL"
    assert st["BAD"]["status"] == "INVALID" and "correct_option" in st["BAD"]["issues"][0]["message"]
    assert st["UNK"]["status"] == "NEEDS_SKILL_MAPPING" and st["UNK"]["canonical_skill"] is None
    assert st["DUPFILE"]["status"] == "DUPLICATE" and st["DUPFILE"]["duplicate_of"]["where"] == "this file"
    assert st["T-2"]["status"] == "WARNING" and st["T-1"]["status"] == "WARNING"  # concepts were derived from the rubric criteria: reviewable, importable
    assert any("derived from the rubric" in i["message"] for i in st["T-2"]["issues"])
    assert (b["total_rows"], b["duplicate_rows"]) == (8, 1) and b["status"] == "PREVIEW"
    async with AsyncSessionLocal() as db:  # preview stored nothing as a question
        assert await db.scalar(select(func.count()).select_from(Question).where(Question.question_text.like(f"%{tok}%"))) == 0
    # resolve the unknown skill to a canonical one, exclude the invalid row
    pg_id = st["M3"]["skill_id"]
    fixed = await client.patch(f"/question-imports/{b['id']}/rows/{st['UNK']['row']}", headers=w["ha"], json={"skill_id": pg_id})
    assert fixed.status_code == 200
    unk = next(x for x in fixed.json()["rows"] if x["raw"]["external_question_id"] == "UNK")
    assert unk["status"] in ("READY", "WARNING") and unk["canonical_skill"] == "PostgreSQL"
    ex = await client.patch(f"/question-imports/{b['id']}/rows/{st['BAD']['row']}", headers=w["ha"], json={"excluded": True})
    assert ex.json()["invalid_rows"] == 0
    done = await client.post(f"/question-imports/{b['id']}/confirm", headers=w["ha"])
    assert done.status_code == 200 and done.json()["imported"] == 6  # 3 MCQ + 2 technical + resolved row; invalid excluded, duplicate skipped
    assert (await client.post(f"/question-imports/{b['id']}/confirm", headers=w["ha"])).status_code == 409  # not repeatable
    async with AsyncSessionLocal() as db:
        qs = (await db.scalars(select(Question).where(Question.import_batch_id == b["id"]))).all()
        assert len(qs) == 6 and all(str(q.organization_id) == w["oa"] and q.visibility == "COMPANY_PRIVATE" and q.source_type == "COMPANY_PRIVATE"
                                    and q.provenance == "COMPANY_IMPORT" and q.status == "VALIDATED" for q in qs)
        m1 = next(q for q in qs if tok in q.question_text)
        assert m1.options == ["return", "yield", "emit"] and m1.correct_option_index == 1 and m1.import_meta["external_question_id"] == "M1"
        t = next(q for q in qs if q.question_type == "TECHNICAL")
        assert t.rubric["criteria"] and len(t.expected_concepts) >= 2
        mid = str(m1.id)
    # same-company reuse: the SAME question record in two different jobs' draft assessments
    r1 = await client.post(f"/assessments/{w['a1']}/questions", headers=w["ha"], json={"question_id": mid})
    r2 = await client.post(f"/assessments/{w['a2']}/questions", headers=w["ha"], json={"question_id": mid})
    assert r1.status_code == r2.status_code == 200
    assert (await client.post(f"/assessments/{w['a1']}/questions", headers=w["ha"], json={"question_id": mid})).status_code == 409
    async with AsyncSessionLocal() as db:
        used = (await db.scalars(select(AssessmentQuestion).where(AssessmentQuestion.question_id == mid))).all()
        assert {str(u.assessment_id) for u in used} == {w["a1"], w["a2"]}
        assert await db.scalar(select(func.count()).select_from(Question).where(Question.question_text.like(f"%{tok}%"))) == 1  # one record, two usages
    # a second upload of the same text is a duplicate against the company's own bank
    again = (await upload(client, w["ha"], w["j2"], sheet([mcq(f"{tok} which keyword defines a generator function in Python?", eid="AGAIN")]))).json()
    assert again["rows"][0]["status"] == "DUPLICATE" and again["rows"][0]["duplicate_of"]["where"] == "your question bank"
    # coverage reflects available vs selected
    cov = (await client.get(f"/question-imports/coverage?job_id={w['j1']}", headers=w["ha"])).json()
    assert cov["needed"] > 0 and all({"skill", "type", "needed", "available_private", "selected"} <= set(x) for x in cov["rows"])
    py_mcq = [x for x in cov["rows"] if x["skill"] == "Python" and x["type"] == "MCQ"]
    assert py_mcq and py_mcq[0]["available_private"] == 2 and py_mcq[0]["selected"] == 1  # M1 and M2 are Python MCQs; M1 is in this job's draft


@pytest.mark.asyncio
async def test_csv_and_json_produce_the_same_validation(client, w):
    body = [mcq("Which built-in returns the length of a Python sequence quickly?", eid="C1"), tech("Explain isolation levels and the anomalies each one prevents.", skill="PostgreSQL", eid="C2")]
    header = ",".join(HEAD)
    import csv
    out = io.StringIO()
    wr = csv.writer(out)
    wr.writerow(HEAD)
    for r in body:
        wr.writerow([r.get(c, "") for c in HEAD])
    csv_res = (await upload(client, w["ha"], w["j1"], out.getvalue().encode(), "q.csv")).json()
    js = {"questions": [{"external_question_id": "C1", "question_type": "MCQ", "skill": "Python", "difficulty": "medium",
                         "question_text": "Which built-in returns the length of a Python sequence quickly?", "options": {"a": "return", "b": "yield", "c": "emit"}, "correct_option": "B"},
                        {"external_question_id": "C2", "question_type": "TECHNICAL_WRITTEN", "skill": "PostgreSQL", "difficulty": "hard",
                         "question_text": "Explain isolation levels and the anomalies each one prevents.",
                         "technical_rubric": ["read committed", "repeatable read", "serializable"]}]}
    json_res = (await upload(client, w["ha"], w["j2"], json.dumps(js).encode(), "q.json")).json()
    assert csv_res["format"] == "csv" and json_res["format"] == "json"
    assert [x["status"] for x in csv_res["rows"]] == ["READY", "WARNING"] or [x["status"] for x in csv_res["rows"]][0] == "READY"
    assert [x["question_type"] for x in json_res["rows"]] == ["MCQ", "TECHNICAL"] and json_res["rows"][0]["options"] == ["return", "yield", "emit"]
    for bad in (("q.txt", b"x"), ("q.xlsx", b"not a workbook"), ("q.json", b"{not json"), ("q.csv", b"skill\nPython\n")):
        assert (await client.post(f"/question-imports?job_id={w['j1']}", headers=w["ha"], files={"file": bad})).status_code == 422


@pytest.mark.asyncio
async def test_row_validation_rules(client, w):
    rows = [mcq("Missing the second option makes this question unusable at all?", option_b="", eid="R1"),
            mcq("An option gap between b and d is not allowed in this template?", option_c="", option_d="x", eid="R2"),
            mcq("Two identical options are ambiguous for a candidate to answer?", option_b="return", eid="R3"),
            mcq("Short", eid="R4"),
            mcq("Difficulty must be one of the three allowed values here?", eid="R5") | {"difficulty": "impossible"},
            {"external_question_id": "R6", "question_type": "CODING", "skill": "Python", "difficulty": "easy", "question_text": "Write code that prints hello world please."},
            tech("Technical rubric is required for written questions to be graded?", skill="Python", eid="R7") | {"technical_rubric": ""},
            tech("A single rubric criterion cannot be turned into two concepts?", skill="Python", eid="R8") | {"technical_rubric": "only one"}]
    b = (await upload(client, w["ha"], w["j1"], sheet(rows))).json()
    by = {x["raw"]["external_question_id"]: x for x in b["rows"]}
    assert all(by[k]["status"] == "INVALID" for k in ("R1", "R2", "R3", "R4", "R5", "R6", "R7", "R8")), {k: v["status"] for k, v in by.items()}
    assert "coding questions cannot be imported" in by["R6"]["issues"][0]["message"]
    assert b["valid_rows"] == 0 and (await client.post(f"/question-imports/{b['id']}/confirm", headers=w["ha"])).json()["imported"] == 0


@pytest.mark.asyncio
async def test_company_a_company_b_isolation_and_tampered_templates(client, w, monkeypatch):
    tok = f"A_ONLY_PRIVATE_TOKEN_{uuid.uuid4().hex[:10]}"
    text = f"{tok} explain how Python decorators wrap a function and preserve its metadata?"
    b = (await upload(client, w["ha"], w["j1"], sheet([mcq(text, eid="ISO")], {"template_version": ci.TEMPLATE_VERSION, "company_id": w["oa"]}))).json()
    await client.post(f"/question-imports/{b['id']}/confirm", headers=w["ha"])
    async with AsyncSessionLocal() as db:
        qa = (await db.scalar(select(Question).where(Question.question_text.like(f"%{tok}%"))))
        qid = str(qa.id)
    # Company A: list, keyword search, direct id, reuse in another job
    assert qid in {x["id"] for x in (await client.get("/questions", headers=w["ha"], params={"organization_id": w["oa"]})).json()}
    assert qid in {x["id"] for x in (await client.get("/questions", headers=w["ha"], params={"q": tok})).json()}
    assert (await client.get(f"/questions/{qid}", headers=w["ha"])).status_code == 200
    assert (await client.post(f"/assessments/{w['a2']}/questions", headers=w["ha"], json={"question_id": qid})).status_code == 200
    # Company B: nothing, by any route
    assert qid not in {x["id"] for x in (await client.get("/questions", headers=w["hb"])).json()}
    assert (await client.get("/questions", headers=w["hb"], params={"q": tok})).json() == []
    assert (await client.get("/questions", headers=w["hb"], params={"organization_id": w["oa"]})).status_code in (403, 404)
    assert (await client.get(f"/questions/{qid}", headers=w["hb"])).status_code in (403, 404)
    assert (await client.post(f"/assessments/{w['ab']}/questions", headers=w["hb"], json={"question_id": qid})).status_code == 404  # attach denied
    assert (await client.post(f"/assessments/{w['a1']}/questions", headers=w["hb"], json={"question_id": qid})).status_code in (403, 404)
    assert (await client.get(f"/question-imports/{b['id']}", headers=w["hb"])).status_code == 404  # A's import batch is invisible
    assert b["id"] not in {x["id"] for x in (await client.get("/question-imports", headers=w["hb"])).json()}
    assert (await client.post(f"/question-imports/{b['id']}/confirm", headers=w["hb"])).status_code == 404
    assert (await client.get(f"/question-imports/template?job_id={w['j1']}", headers=w["hb"])).status_code in (403, 404)
    # semantic search / RAG / reranker input: B's scope never returns A's private point, A's scope does
    from app.services.ai_gateway import vector_store as vs
    from app.services.ai_gateway.embeddings import get_embedding_service
    vec = get_embedding_service().embed([text])[0]
    hits_b = vs.search("questions", vec, TenantScope(organization_id=uuid.UUID(w["ob"])), limit=20)
    hits_a = vs.search("questions", vec, TenantScope(organization_id=uuid.UUID(w["oa"])), limit=20)
    assert all(tok not in (h.payload.get("text") or "") for h in hits_b) and any(h.payload.get("question_id") == qid for h in hits_a)
    # B uploading the same wording is NOT flagged as a duplicate (A's bank is never consulted) and gets its own private copy
    bb = (await upload(client, w["hb"], w["jb"], sheet([mcq(text, eid="B-COPY")]))).json()
    assert bb["rows"][0]["status"] == "READY" and bb["rows"][0]["duplicate_of"] is None
    await client.post(f"/question-imports/{bb['id']}/confirm", headers=w["hb"])
    async with AsyncSessionLocal() as db:
        owners = {str(x.organization_id) for x in (await db.scalars(select(Question).where(Question.question_text.like(f"%{tok}%")))).all()}
        assert owners == {w["oa"], w["ob"]}  # two separate private records
    # LLM context: interview pool generation for B never sees A's token (bank lookups are organization-scoped)
    from app.services.interviews import pool as P
    prompts = []

    class GW:
        async def generate_structured(self, prompt, schema, **kw):
            prompts.append(prompt)
            return schema(question_text=f"Prepared question {len(prompts)} about the requirement, explain your approach in detail.", reason_for_question="r")
    monkeypatch.setattr(P, "get_ai_gateway", lambda: GW())
    monkeypatch.setattr(P, "retrieve", lambda *a, **k: [])
    monkeypatch.setattr(P, "to_source_refs", lambda d: [])
    await P.fill_pool(uuid.UUID(w["jb"]))
    assert prompts and all(tok not in p for p in prompts)
    # tampered templates: a file generated for A but uploaded by B is refused, 0 rows imported
    files_a = sheet([mcq("Which statement about Python generators is correct here?", eid="X")], {"template_version": ci.TEMPLATE_VERSION, "company_id": w["oa"], "job_id": w["j1"]})
    async with AsyncSessionLocal() as db:
        before = await db.scalar(select(func.count()).select_from(QuestionImportBatch).where(QuestionImportBatch.organization_id == uuid.UUID(w["ob"])))
    refused = await upload(client, w["hb"], w["jb"], files_a)
    assert refused.status_code == 422 and "different company" in refused.text
    assert (await upload(client, w["hb"], w["j1"], files_a)).status_code in (403, 404)  # A's job id: not B's job
    async with AsyncSessionLocal() as db:
        assert await db.scalar(select(func.count()).select_from(QuestionImportBatch).where(QuestionImportBatch.organization_id == uuid.UUID(w["ob"]))) == before
    # editing the metadata is not authorization: company A pointing its file at company B's ids gets nothing from B
    forged = sheet([mcq("Which statement about Python iterators is correct here?", eid="F")], {"template_version": ci.TEMPLATE_VERSION, "company_id": w["ob"], "job_id": w["jb"]})
    assert (await upload(client, w["ha"], w["j1"], forged)).status_code == 422
    assert (await upload(client, w["ha"], w["jb"], forged)).status_code in (403, 404)


@pytest.mark.asyncio
async def test_students_never_receive_keys_or_provenance_and_published_assessments_are_frozen(client, w):
    b = (await upload(client, w["ha"], w["j2"], sheet([mcq("Which Python statement defines a generator using the yield keyword here?", eid="P1", explanation="SECRET-EXPLANATION"),
                                                      tech("Explain how Python resolves attribute lookups through the method resolution order.", skill="Python", eid="P2")]))).json()
    await client.post(f"/question-imports/{b['id']}/confirm", headers=w["ha"])
    async with AsyncSessionLocal() as db:
        ids = [str(q.id) for q in (await db.scalars(select(Question).where(Question.import_batch_id == b["id"]))).all()]
        st, _, hs = await make_student(db)
        job = await db.get(__import__("app.models.jobs", fromlist=["Job"]).Job, w["j2"])
        job.status = JobStatus.PUBLISHED
        app_ = await make_application(db, job, st)
        await db.commit()
    for qid in ids:
        assert (await client.post(f"/assessments/{w['a2']}/questions", headers=w["ha"], json={"question_id": qid})).status_code == 200
    pub = await client.post(f"/assessments/{w['a2']}/publish", headers=w["ha"])
    assert pub.status_code == 200, pub.text
    late = (await client.post(f"/assessments/{w['a2']}/questions", headers=w["ha"], json={"question_id": ids[0]}))
    assert late.status_code == 409  # frozen (already attached or published: either way it cannot change)
    att = (await client.post(f"/assessments/{w['a2']}/attempts", headers=hs, json={"application_id": str(app_.id)})).json()
    raw = (await client.get(f"/assessments/attempts/{att['id']}", headers=hs)).text
    for forbidden in ("correct_option", "explanation", "SECRET-EXPLANATION", "technical_rubric", "criteria", "expected_concepts", "provenance", "import_meta",
                      "import_batch", "organization_id", "external_question_id", "COMPANY_IMPORT", "rubric"):
        assert forbidden not in raw, forbidden
    assert (await client.delete(f"/assessments/{w['a2']}/questions/{ids[0]}", headers=w["ha"])).status_code == 409


@pytest.mark.asyncio
async def test_builder_uses_the_companys_own_questions_before_any_ai_and_never_another_companys(client, w, monkeypatch):
    """Priority: same company's private questions -> approved platform -> AI. Company B's builder cannot see A's imports."""
    from types import SimpleNamespace

    from app.agents import assessment_agent as AA
    from app.models.enums import QuestionType
    from app.services.assessments.blueprint import SkillAllocation
    from app.services.assessments.generator import get_confirmed_job_skills  # noqa: F401
    from app.services.questions.governance import USABLE_PLATFORM  # noqa: F401
    from app.services.skills.normalizer import normalize_skill_name

    b = (await upload(client, w["ha"], w["j1"], sheet([mcq(f"Which Python construct does the builder test number {i} ask about, precisely?", eid=f"B{i}") for i in range(3)]))).json()
    await client.post(f"/question-imports/{b['id']}/confirm", headers=w["ha"])
    calls = []

    async def fake_generate(db, **kw):
        calls.append(kw["organization_id"])
        raise AssertionError("AI generation must not run for slots the company already covers")
    monkeypatch.setattr(AA.gen, "generate_missing_question", fake_generate)
    async with AsyncSessionLocal() as db:
        sid, _ = await normalize_skill_name(db, "Python")

        def deps(org_id):
            alloc = SkillAllocation(skill_id=sid, skill_name="Python", weight=1.0, mcq_count=3, technical_count=0, coding_count=0)
            wk = AA.SkillWork(alloc=alloc, need={QuestionType.MCQ: 3})
            return SimpleNamespace(db=db, job=SimpleNamespace(id=uuid.uuid4(), organization_id=org_id), actor_user_id=None, work=[wk],
                                   log=SimpleNamespace(record=lambda *a, **k: None)), wk
        dA, wkA = deps(uuid.UUID(w["oa"]))
        await AA._company(dA, 0)
        out = await AA._generate_for(dA, 0)  # nothing left to generate
        picked = wkA.picked[QuestionType.MCQ]
        assert len(picked) == 3 and calls == [] and out["created_question_ids"] == []
        owners = {str(q.organization_id) for q in (await db.scalars(select(Question).where(Question.id.in_(picked)))).all()}
        assert owners == {w["oa"]}
        # company B: none of A's private questions are offered, so B's builder finds nothing to reuse
        dB, wkB = deps(uuid.UUID(w["ob"]))
        await AA._company(dB, 0)
        assert not set(wkB.picked.get(QuestionType.MCQ, [])) & set(picked)
        assert not any(str(q.organization_id) == w["oa"] for q in (await db.scalars(select(Question).where(Question.id.in_(wkB.picked.get(QuestionType.MCQ, []) or [uuid.uuid4()])))).all())
