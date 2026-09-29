"""Domain-aware question bank and import, stage language narrowing, per-stage proctoring, per-stage analytics, student stage hint."""
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.database import AsyncSessionLocal
from app.models.enums import JobStatus
from app.models.pipeline import HiringStage
from app.models.questions import Question
from app.services.assessments import versioning as ver
from app.services.pipeline import stages as S
from app.services.questions import company_import as ci
from tests.factories import configure_pipeline, make_application, make_company, make_job, make_stage_assessment, make_student
from tests.integration.test_hiring_pipeline import SKILLS


async def _company(name):
    async with AsyncSessionLocal() as db:
        org, rec, hc = await make_company(db, name)
        job = await make_job(db, org, rec, SKILLS, status=JobStatus.REQUIREMENTS_CONFIRMED)
        await db.commit()
        return org, hc, job


def _workbook(rows: list[dict], meta: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(ci.COLUMNS)
    for r in rows:
        ws.append([r.get(c, "") for c in ci.COLUMNS])
    m = wb.create_sheet("Metadata")
    m.append(["key", "value"])
    for k, v in meta.items():
        m.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


APT = dict(question_type="MCQ", assessment_domain="APTITUDE", difficulty="medium", option_a="40", option_b="42", option_c="44", option_d="48", correct_option="B")


@pytest.mark.asyncio
async def test_domain_templates_carry_domain_columns_and_example_rows(client):
    org, hc, job = await _company("TplCo")
    for domain, expects in (("APTITUDE", "Quantitative Aptitude"), ("HR_INTERVIEW", "communication"), ("TECHNICAL", None)):
        r = await client.get(f"/question-imports/template?job_id={job.id}&format=xlsx&domain={domain}", headers=hc)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        header = [c.value for c in wb["Questions"][1]]
        assert header[-3:] == ["assessment_domain", "category", "sub_category"] and header[:2] == ["external_question_id", "question_type"]
        meta = {row[0].value: row[1].value for row in wb["Metadata"].iter_rows(min_row=2)}
        assert meta["assessment_domain"] == domain and meta["company_id"] == str(org.id)
        if expects:
            assert any(expects in str(c.value) for row in wb["Questions"].iter_rows(min_row=2) for c in row)  # a filled example row
    assert (await client.get(f"/question-imports/template?job_id={job.id}&domain=NOPE", headers=hc)).status_code == 422
    js = (await client.get(f"/question-imports/template?job_id={job.id}&format=json&domain=APTITUDE", headers=hc)).json()
    assert js["questions"][0]["category"] == "Quantitative Aptitude" and js["questions"][0]["assessment_domain"] == "APTITUDE"


@pytest.mark.asyncio
async def test_aptitude_and_hr_import_uses_categories_validates_and_never_imports_examples(client):
    org, hc, job = await _company("ImpCo")
    meta = {"template_version": ci.TEMPLATE_VERSION, "company_id": str(org.id), "job_id": str(job.id), "assessment_domain": "APTITUDE"}
    rows = [
        {**APT, "category": "quantitative aptitude", "question_text": "Import test: what is 6 multiplied by 7 in this exercise?"},           # ok, category normalized
        {**APT, "category": "Logical Reasoning", "question_text": "Import test: which number comes next in the series 2, 4, 8, 16?", "correct_option": "A", "option_a": "32"},
        {**APT, "category": "", "question_text": "Import test: this row has no category at all in it."},                                  # invalid
        {**APT, "category": "Verbal Ability", "question_type": "TECHNICAL_WRITTEN", "question_text": "Import test: a written aptitude question is not supported."},  # invalid
        {**APT, "category": "Quantitative Aptitude", "question_text": "Import test: what is 6 multiplied by 7 in this exercise?"},        # duplicate in file
        {"question_type": "TECHNICAL_WRITTEN", "assessment_domain": "HR_INTERVIEW", "category": "communication",
         "question_text": "Tell me about a time you explained something complex to a new teammate."},                                   # ok
        {"question_type": "TECHNICAL_WRITTEN", "assessment_domain": "HR_INTERVIEW", "category": "motivation",
         "question_text": "Are you planning to have children in the next few years?"},                                                    # sensitive -> invalid
        {"question_type": "TECHNICAL_WRITTEN", "assessment_domain": "HR_INTERVIEW", "category": "astrology", "question_text": "What is your favourite programming language and why?"},  # bad category
        {**APT, "category": "Quantitative Aptitude", "question_text": "EXAMPLE (delete this row): sample text that must never be imported."},
    ]
    up = await client.post(f"/question-imports?job_id={job.id}&domain=APTITUDE", headers=hc,
                           files={"file": ("q.xlsx", _workbook(rows, meta), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert up.status_code == 200, up.text
    b = up.json()
    by = {r["row"]: r for r in b["rows"]}
    assert len(b["rows"]) == 8  # the example row was dropped
    assert by[2]["status"] in ("READY", "WARNING") and by[2]["category"] == "Quantitative Aptitude" and by[2]["skill_id"] is None and by[2]["domain"] == "APTITUDE"
    assert by[3]["status"] in ("READY", "WARNING")
    assert by[4]["status"] == "INVALID" and any("category" in i["message"] for i in by[4]["issues"])
    assert by[5]["status"] == "INVALID" and any("MCQ" in i["message"] for i in by[5]["issues"])
    assert by[6]["status"] == "DUPLICATE"
    assert by[7]["status"] in ("READY", "WARNING") and by[7]["domain"] == "HR_INTERVIEW" and by[7]["category"] == "communication"
    assert by[8]["status"] == "INVALID" and any("sensitive" in i["message"] for i in by[8]["issues"])
    assert by[9]["status"] == "INVALID" and any("category" in i["message"] for i in by[9]["issues"])
    assert not any(r["status"] == "NEEDS_SKILL_MAPPING" for r in b["rows"])  # aptitude / HR rows never ask for a skill
    ok = await client.post(f"/question-imports/{b['id']}/confirm", headers=hc)
    assert ok.status_code == 200 and ok.json()["imported"] == 3
    async with AsyncSessionLocal() as db:
        made = (await db.scalars(select(Question).where(Question.organization_id == org.id, Question.import_batch_id == b["id"]))).all()
        assert {(q.domain, q.category, q.skill_id, q.visibility, q.provenance) for q in made} == {
            ("APTITUDE", "Quantitative Aptitude", None, "COMPANY_PRIVATE", "COMPANY_IMPORT"), ("APTITUDE", "Logical Reasoning", None, "COMPANY_PRIVATE", "COMPANY_IMPORT"),
            ("HR_INTERVIEW", "communication", None, "COMPANY_PRIVATE", "COMPANY_IMPORT")}
        assert not any("EXAMPLE" in q.question_text for q in made)
    # the bank can be filtered by domain and category, and another company sees none of it
    lst = (await client.get(f"/questions?organization_id={org.id}&domain=APTITUDE", headers=hc)).json()
    assert {q["category"] for q in lst} >= {"Quantitative Aptitude", "Logical Reasoning"} and all(q["domain"] == "APTITUDE" and q["skill_id"] is None for q in lst)
    assert [q["category"] for q in (await client.get(f"/questions?organization_id={org.id}&domain=HR_INTERVIEW", headers=hc)).json()] == ["communication"]
    async with AsyncSessionLocal() as db:
        _, _, other = await make_company(db, "ImpOther")
        await db.commit()
    assert "Import test" not in (await client.get("/questions", headers=other)).text
    # coverage for the aptitude stage: needed slots by category against what the company already has
    async with AsyncSessionLocal() as db:
        await configure_pipeline(db, await db.get(type(job), job.id), [S.APTITUDE])
        st = await db.scalar(select(HiringStage).where(HiringStage.job_id == job.id, HiringStage.stage_type == S.APTITUDE))
        st.question_count, st.config = 4, {"categories": {"Quantitative Aptitude": 50, "Logical Reasoning": 50}, "difficulty": {"easy": 0, "medium": 100, "hard": 0}}
        await db.commit()
    cov = (await client.get(f"/question-imports/coverage?job_id={job.id}&domain=APTITUDE", headers=hc)).json()
    assert cov["needed"] == 4 and {r["category"]: r["available_private"] for r in cov["rows"]} == {"Quantitative Aptitude": 1, "Logical Reasoning": 1}
    assert cov["covered"] == 2 and cov["gap"] == 2


@pytest.mark.asyncio
async def test_manual_creation_supports_aptitude_and_rejects_unsafe_hr_questions(client):
    org, hc, job = await _company("ManCo")
    good = await client.post("/questions", headers=hc, json={"question_text": "Manual test: what is 9 multiplied by 9 here?", "question_type": "MCQ", "domain": "APTITUDE",
                                                             "category": "Quantitative Aptitude", "difficulty": "easy", "options": ["72", "81", "90", "99"], "correct_option": 1})
    assert good.status_code == 200 and good.json()["created"] == 1
    tech_without_skill = await client.post("/questions", headers=hc, json={"question_text": "Manual test: explain how a hash map handles collisions.", "question_type": "TECHNICAL"})
    assert tech_without_skill.json()["created"] == 0 and "skill is required" in tech_without_skill.json()["errors"][0]["error"]
    bad = await client.post("/questions", headers=hc, json={"question_text": "Manual test: what is your religion and marital status?", "question_type": "TECHNICAL",
                                                            "domain": "HR_INTERVIEW", "category": "motivation"})
    assert bad.json()["created"] == 0 and "sensitive" in bad.json()["errors"][0]["error"]
    fine = await client.post("/questions", headers=hc, json={"question_text": "Manual test: what kind of team environment helps you do your best work?", "question_type": "TECHNICAL",
                                                             "domain": "HR_INTERVIEW", "category": "work preferences"})
    assert fine.json()["created"] == 1


# ------------------------------------------------------------------ language narrowing, proctoring, analytics, hint
@pytest.mark.asyncio
async def test_coding_stage_languages_narrow_the_frozen_problems_without_touching_the_shared_question(client):
    async with AsyncSessionLocal() as db:
        org, rec, hc = await make_company(db, "LangCo")
        job = await make_job(db, org, rec, SKILLS, status=JobStatus.PUBLISHED)
        await configure_pipeline(db, job, [S.CODING])
        st, a, aqs, qs = await make_stage_assessment(db, org, job, S.CODING, published=False)
        a.config = {**a.config, "allowed_languages": ["python"]}
        await ver.ensure_version(db, a)
        st.status, a.status = "PUBLISHED", "PUBLISHED"
        stu, _, hs = await make_student(db)
        app_ = await make_application(db, job, stu)
        qid = qs[0].id
        await db.commit()
    att = (await client.post(f"/assessments/{a.id}/attempts", headers=hs, json={"application_id": str(app_.id)})).json()["id"]
    sess = (await client.get(f"/assessments/attempts/{att}", headers=hs)).json()
    coding = [q["question"] for s in sess["sections"] for q in s["questions"]][0]
    assert coding["allowed_languages"] == ["python"]  # the recruiter's choice, enforced on the frozen problem
    assert "test_cases" not in str(sess) and coding["hidden_test_count"] >= 1  # hidden tests stay server-side
    async with AsyncSessionLocal() as db:
        assert (await db.get(Question, qid)).allowed_languages is None  # the shared question row is untouched


@pytest.mark.asyncio
async def test_proctoring_can_be_switched_off_for_a_stage_and_stays_enforced_for_the_others(client, monkeypatch):
    from app.api.v1 import proctoring as pr

    monkeypatch.setattr(pr.settings, "PROCTOR_ENFORCE", True)
    async with AsyncSessionLocal() as db:
        org, rec, hc = await make_company(db, "ProcCo")
        job = await make_job(db, org, rec, SKILLS, status=JobStatus.PUBLISHED)
        await configure_pipeline(db, job, [S.APTITUDE, S.TECHNICAL])
        apt = await make_stage_assessment(db, org, job, S.APTITUDE)
        tech = await make_stage_assessment(db, org, job, S.TECHNICAL)
        apt[0].proctored, tech[0].proctored = False, True
        stu, _, hs = await make_student(db)
        app_ = await make_application(db, job, stu)
        await db.commit()
    ok = await client.post(f"/assessments/{apt[1].id}/attempts", headers=hs, json={"application_id": str(app_.id)})
    assert ok.status_code == 200  # this stage is not proctored: no session needed
    assert (await client.post(f"/assessments/attempts/{ok.json()['id']}/submit", headers=hs)).json()["completed"]
    blocked = await client.post(f"/assessments/{tech[1].id}/attempts", headers=hs, json={"application_id": str(app_.id)})
    assert blocked.status_code == 428 and blocked.json()["detail"]["code"] == "PROCTORING_REQUIRED"


@pytest.mark.asyncio
async def test_stage_analytics_and_the_students_next_stage_hint(client):
    async with AsyncSessionLocal() as db:
        org, rec, hc = await make_company(db, "AnaCo")
        job = await make_job(db, org, rec, SKILLS, status=JobStatus.PUBLISHED)
        await configure_pipeline(db, job, [S.APTITUDE, S.TECHNICAL, S.HR_INTERVIEW])
        apt = await make_stage_assessment(db, org, job, S.APTITUDE)
        await make_stage_assessment(db, org, job, S.TECHNICAL)
        stu, _, hs = await make_student(db)
        app_ = await make_application(db, job, stu)
        await db.commit()
    mine = (await client.get("/applications/mine", headers=hs)).json()[0]
    assert (mine["next_stage_label"], mine["next_stage_status"], mine["stages_total"], mine["stages_done"]) == ("Aptitude Assessment", "AVAILABLE", 3, 0)
    att = (await client.post(f"/assessments/{apt[1].id}/attempts", headers=hs, json={"application_id": str(app_.id)})).json()["id"]
    q = (await client.get(f"/assessments/attempts/{att}", headers=hs)).json()["sections"][0]["questions"][0]
    await client.put(f"/assessments/attempts/{att}/answers", headers=hs, json={"assessment_question_id": q["id"], "selected_option_index": 1})
    await client.post(f"/assessments/attempts/{att}/submit", headers=hs)
    mine = (await client.get("/applications/mine", headers=hs)).json()[0]
    assert (mine["next_stage_label"], mine["stages_done"]) == ("Technical Assessment", 1)
    ana = (await client.get(f"/hiring-pipeline/jobs/{job.id}/analytics", headers=hc)).json()
    by = {s["stage_type"]: s for s in ana["stages"]}
    assert list(by) == [S.APTITUDE, S.TECHNICAL, S.HR_INTERVIEW]
    assert by[S.APTITUDE]["completed"] == 1 and by[S.APTITUDE]["average_score_pct"] == pytest.approx(33.3, abs=0.1) and by[S.TECHNICAL]["waiting"] == 1
    assert by[S.TECHNICAL]["candidates_reached"] == 1 and by[S.HR_INTERVIEW]["interviews_completed"] == 0
    assert not any("overall" in k for s in ana["stages"] for k in s)  # per-stage numbers only, no blended score
    _, _, other = (await _other())
    assert (await client.get(f"/hiring-pipeline/jobs/{job.id}/analytics", headers=other)).status_code in (403, 404)


async def _other():
    async with AsyncSessionLocal() as db:
        r = await make_company(db, "AnaOther")
        await db.commit()
        return r


@pytest.mark.asyncio
async def test_a_structured_interview_cannot_be_ended_after_one_question(client, monkeypatch):
    from app.api.v1 import interviews as iv_api
    from app.schemas.rubric import RubricEvaluation
    from app.services.interviews import pool as P
    from tests.integration.test_hiring_pipeline import FakeGW

    gw = FakeGW()
    monkeypatch.setattr(P, "get_ai_gateway", lambda: gw)
    monkeypatch.setattr(P, "retrieve", lambda *a, **k: [])
    monkeypatch.setattr(P, "to_source_refs", lambda docs: [])

    async def ev(turn):
        return RubricEvaluation(concept_accuracy=0.8, reasoning=0.8, completeness=0.8, communication=0.8, demonstrated_concepts=[], missing_concepts=[], evaluator_confidence=0.9)
    monkeypatch.setattr(iv_api, "evaluate_turn_answer", ev)
    async with AsyncSessionLocal() as db:
        org, rec, hc = await make_company(db, "ShortCo")
        job = await make_job(db, org, rec, SKILLS, status=JobStatus.PUBLISHED)
        await configure_pipeline(db, job, [S.TECH_INTERVIEW])
        stu, _, hs = await make_student(db)
        app_ = await make_application(db, job, stu)
        await db.commit()
    await P.fill_pool(job.id)
    iv = (await client.post("/interviews/start", headers=hs, json={"application_id": str(app_.id)})).json()
    assert iv["question_budget"] >= 6 and iv["min_questions"] >= 6 and "plan" not in iv
    t = (await client.post(f"/interviews/{iv['id']}/next-turn", headers=hs)).json()
    await client.post(f"/interviews/turns/{t['id']}/answer", headers=hs, json={"answer_text": "A single detailed answer. " * 5})
    early = await client.post(f"/interviews/{iv['id']}/finish", headers=hs)
    assert early.status_code == 409 and early.json()["detail"]["code"] == "INTERVIEW_TOO_SHORT" and "at least" in early.json()["detail"]["message"]


@pytest.mark.asyncio
async def test_technical_interview_length_follows_its_duration_when_no_count_is_given(client):
    """Regression (found in the browser): changing the interview length sent an empty question count, which the server ignored,
    so a '30 minute' interview kept the old 8-question length."""
    from app.services.interviews import pool as P

    async with AsyncSessionLocal() as db:
        org, rec, hc = await make_company(db, "LenCo")
        job = await make_job(db, org, rec, SKILLS, status=JobStatus.REQUIREMENTS_CONFIRMED)
        await db.commit()
    v = (await client.get(f"/hiring-pipeline/jobs/{job.id}", headers=hc)).json()
    ti = next(s for s in v["stages"] if s["stage_type"] == S.TECH_INTERVIEW)
    assert ti["question_count"] is None  # the default no longer pins a length
    for minutes, expect in ((30, 6), (45, 9), (60, 10)):
        body = {"stages": [{"stage_type": S.TECH_INTERVIEW, "enabled": True, "duration_minutes": minutes, "question_count": None,
                            "config": {"min_questions": 6, "max_questions": 10}}]}
        r = await client.put(f"/hiring-pipeline/jobs/{job.id}", headers=hc, json=body)
        assert r.status_code == 200, r.text
        async with AsyncSessionLocal() as db:
            from app.models.jobs import Job

            t = await P.upsert_template(db, await db.get(Job, job.id), S.TECH_INTERVIEW)
            assert t.config["question_budget"] == expect and t.config["recommended_minutes"] == minutes
            await db.commit()
