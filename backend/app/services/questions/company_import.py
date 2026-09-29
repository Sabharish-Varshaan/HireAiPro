"""Company-private question import: template -> parse -> validate -> skill mapping -> duplicates -> preview -> confirm.

Everything is scoped to ONE organization. Duplicate detection only ever looks at that organization's private bank, its current assessment
and the public platform bank (semantic search uses TenantScope(organization_id=that org)); another company's private content is never
queried. Nothing becomes a Question until the recruiter confirms the batch.
"""

import asyncio
import csv
import datetime as dt
import io
import json
import re
import uuid
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.assessments import AssessmentQuestion
from app.models.enums import QuestionSourceType, QuestionStatus as QS, QuestionType, Visibility
from app.models.questions import Question, QuestionImportBatch
from app.models.skills import Skill
from app.services.ai_gateway.vector_store import TenantScope
from app.services.questions.validator import content_hash, index_question, validate_semantics, validate_structure
from app.services.skills.normalizer import normalize_skill_name

TEMPLATE_VERSION = "company-questions-v1"
MAX_ROWS = 500
MAX_BYTES = 2_000_000
COLUMNS = ["external_question_id", "question_type", "skill", "difficulty", "question_text", "option_a", "option_b", "option_c", "option_d", "option_e",
           "correct_option", "explanation", "technical_rubric", "expected_concepts", "max_score", "tags", "time_limit_seconds", "required"]
REQUIRED_MCQ = ["skill", "difficulty", "question_text", "option_a", "option_b", "correct_option"]
REQUIRED_TECH = ["external_question_id", "skill", "difficulty", "question_text", "technical_rubric"]
LETTERS = "ABCDE"
DIFFICULTIES = ("easy", "medium", "hard")
TYPE_ALIASES = {"MCQ": "MCQ", "MULTIPLE_CHOICE": "MCQ", "TECHNICAL_WRITTEN": "TECHNICAL", "TECHNICAL": "TECHNICAL", "WRITTEN": "TECHNICAL"}
STATUS_ORDER = ["INVALID", "NEEDS_SKILL_MAPPING", "DUPLICATE", "WARNING", "READY"]


class ImportFileError(ValueError):
    """The file itself cannot be used (format, size, wrong company). Nothing is stored."""


# ---------------------------------------------------------------- templates
def _instructions() -> list[list[str]]:
    return [
        ["HireAiPro company question template"],
        [""],
        ["Add one question per row on the 'Questions' sheet. Do not rename or reorder the columns. Do not edit the 'Metadata' sheet."],
        ["question_type: MCQ or TECHNICAL_WRITTEN. Coding questions cannot be imported here."],
        ["MCQ required: skill, difficulty, question_text, option_a, option_b, correct_option (a letter A-E that matches an option you filled)."],
        ["TECHNICAL_WRITTEN required: external_question_id, skill, difficulty, question_text, technical_rubric (grading criteria separated by |)."],
        ["Optional: option_c-option_e, explanation, expected_concepts (separated by |), max_score, tags (separated by |), time_limit_seconds, required (yes/no)."],
        ["difficulty: easy, medium or hard. skill: a skill name or common alias (for example Postgres or python3); unknown names are mapped during the preview."],
        ["These questions become PRIVATE to your company. Other companies can never see, search or reuse them. You review a preview before anything is imported."],
        [""],
        ["Example MCQ:  MCQ | Python | medium | Which keyword defines a generator function? | return | yield | emit | | | B"],
        ["Example written:  TECHNICAL_WRITTEN | PostgreSQL | hard | Explain when a B-tree index is chosen by the planner. | criteria 1 | criteria 2"],
    ]


def build_metadata(*, org, job, assessment) -> dict:
    return {"template_version": TEMPLATE_VERSION, "template_id": str(uuid.uuid4()), "company_id": str(org.id), "company_name": org.name,
            "job_id": str(job.id), "job_title": job.title, "assessment_id": str(assessment.id) if assessment else "",
            "assessment_name": assessment.title if assessment else "", "generated_at": dt.datetime.now(dt.timezone.utc).isoformat()}


def build_xlsx(meta: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    for row in _instructions():
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 130
    q = wb.create_sheet("Questions")
    q.append(COLUMNS)
    for c in q[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
    q.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJKLMNOPQR", (18, 18, 20, 12, 60, 28, 28, 28, 28, 28, 14, 30, 40, 30, 10, 20, 12, 10)):
        q.column_dimensions[col].width = w
    for col, values in (("B", '"MCQ,TECHNICAL_WRITTEN"'), ("D", '"easy,medium,hard"'), ("K", '"A,B,C,D,E"'), ("R", '"yes,no"')):
        dv = DataValidation(type="list", formula1=values, allow_blank=True)
        q.add_data_validation(dv)
        dv.add(f"{col}2:{col}{MAX_ROWS + 1}")
    m = wb.create_sheet("Metadata")
    m.append(["key", "value"])
    for k, v in meta.items():
        m.append([k, v])
    m.column_dimensions["A"].width = 22
    m.column_dimensions["B"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(meta: dict) -> bytes:
    out = io.StringIO()
    for k, v in meta.items():
        out.write(f"# {k}={v}\n")
    w = csv.writer(out)
    w.writerow(COLUMNS)
    return out.getvalue().encode("utf-8")


def build_json(meta: dict) -> bytes:
    example_mcq = {"external_question_id": "EX-1", "question_type": "MCQ", "skill": "Python", "difficulty": "medium",
                   "question_text": "Which keyword defines a generator function?", "options": {"a": "return", "b": "yield", "c": "emit"}, "correct_option": "B",
                   "explanation": "yield turns the function into a generator", "tags": ["basics"], "time_limit_seconds": 60, "required": False}
    example_tech = {"external_question_id": "EX-2", "question_type": "TECHNICAL_WRITTEN", "skill": "PostgreSQL", "difficulty": "hard",
                    "question_text": "Explain when the planner chooses a B-tree index scan over a sequential scan.",
                    "technical_rubric": ["selectivity and statistics", "cost estimates", "index-only scans"], "expected_concepts": ["selectivity", "statistics"], "max_score": 5}
    doc = {"metadata": meta, "instructions": "Replace the examples. correct_option is a letter A-E matching a filled option. Coding questions are not supported.",
           "questions": [example_mcq, example_tech]}
    return json.dumps(doc, indent=2).encode("utf-8")


# ---------------------------------------------------------------- parsing
def _norm_key(k: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(k or "").strip().lower()).strip("_")


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _row_from_flat(d: dict) -> dict:
    return {c: _cell(d.get(c)) for c in COLUMNS}


def parse_upload(raw: bytes, filename: str) -> tuple[list[dict], dict, str]:
    """Returns (rows, metadata, format). Raises ImportFileError for anything unusable."""
    if len(raw) > MAX_BYTES:
        raise ImportFileError("File is larger than 2 MB")
    name = (filename or "").lower()
    try:
        if name.endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            if "Questions" not in wb.sheetnames:
                raise ImportFileError("The workbook has no 'Questions' sheet (use the downloaded template)")
            meta = {}
            if "Metadata" in wb.sheetnames:
                for r in list(wb["Metadata"].iter_rows(values_only=True))[1:]:
                    if r and r[0]:
                        meta[_norm_key(r[0])] = _cell(r[1] if len(r) > 1 else "")
            it = wb["Questions"].iter_rows(values_only=True)
            header = [_norm_key(h) for h in next(it, [])]
            rows = [dict(zip(header, r)) for r in it if r and any(_cell(c) for c in r)]
            fmt = "xlsx"
        elif name.endswith(".csv"):
            text = raw.decode("utf-8-sig")
            meta, lines = {}, []
            for line in text.splitlines():
                if line.startswith("#"):
                    k, _, v = line[1:].strip().partition("=")
                    meta[_norm_key(k)] = v.strip()
                else:
                    lines.append(line)
            reader = csv.DictReader(io.StringIO("\n".join(lines)))
            rows = [{_norm_key(k): v for k, v in r.items() if k} for r in reader if any((v or "").strip() for v in r.values())]
            fmt = "csv"
        elif name.endswith(".json"):
            data = json.loads(raw.decode("utf-8"))
            meta = {_norm_key(k): _cell(v) for k, v in (data.get("metadata") or {}).items()} if isinstance(data, dict) else {}
            items = data.get("questions", []) if isinstance(data, dict) else data
            rows = []
            for it_ in items:
                it_ = {_norm_key(k): v for k, v in it_.items()}
                opts = it_.pop("options", None)
                if isinstance(opts, dict):
                    for k, v in opts.items():
                        it_[f"option_{str(k).lower()}"] = v
                elif isinstance(opts, list):
                    for i, v in enumerate(opts[:5]):
                        it_[f"option_{'abcde'[i]}"] = v
                for lk in ("technical_rubric", "expected_concepts", "tags"):
                    if isinstance(it_.get(lk), list):
                        it_[lk] = "|".join(str(x) for x in it_[lk])
                rows.append(it_)
            fmt = "json"
        else:
            raise ImportFileError("Upload an .xlsx, .csv or .json file")
    except ImportFileError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError(f"Could not read the file: {type(exc).__name__}") from exc
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Too many rows (max {MAX_ROWS})")
    if not rows:
        raise ImportFileError("The file has no question rows")
    missing_cols = {"question_text", "skill"} - {k for r in rows for k in r}
    if missing_cols:
        raise ImportFileError(f"Missing required column(s): {', '.join(sorted(missing_cols))}")
    return [_row_from_flat(r) for r in rows], meta, fmt


# ---------------------------------------------------------------- validation
@dataclass
class RowResult:
    row: int
    status: str = "READY"
    issues: list[dict] = field(default_factory=list)
    raw: dict = field(default_factory=dict)
    question_type: str | None = None
    skill_id: str | None = None
    canonical_skill: str | None = None
    difficulty: str | None = None
    options: list[str] | None = None
    correct_index: int | None = None
    criteria: list[str] | None = None
    concepts: list[str] | None = None
    duplicate_of: dict | None = None
    excluded: bool = False

    def bad(self, msg: str, level: str = "error") -> None:
        self.issues.append({"level": level, "message": msg})

    def as_dict(self) -> dict:
        return {"row": self.row, "status": self.status, "issues": self.issues, "raw": self.raw, "question_type": self.question_type,
                "skill_id": self.skill_id, "canonical_skill": self.canonical_skill, "difficulty": self.difficulty, "options": self.options,
                "correct_index": self.correct_index, "criteria": self.criteria, "concepts": self.concepts, "duplicate_of": self.duplicate_of,
                "excluded": self.excluded}


def _split(s: str) -> list[str]:
    return [p.strip() for p in re.split(r"[|\n]", s or "") if p.strip()]


def _truthy(s: str) -> bool:
    return (s or "").strip().lower() in ("yes", "y", "true", "1", "required")


def _norm_text(t: str) -> str:
    return " ".join((t or "").lower().split())


def check_structure(r: RowResult) -> None:
    """Schema-level validation of one row (no database)."""
    raw = r.raw
    qt = TYPE_ALIASES.get(re.sub(r"[\s-]+", "_", raw["question_type"].strip().upper()))
    if qt is None:
        r.bad("question_type must be MCQ or TECHNICAL_WRITTEN" + (" (coding questions cannot be imported)" if raw["question_type"].strip().upper() == "CODING" else ""))
        return
    r.question_type = qt
    for col in (REQUIRED_MCQ if qt == "MCQ" else REQUIRED_TECH):
        if not raw[col]:
            r.bad(f"{col} is required")
    d = raw["difficulty"].strip().lower()
    if d and d not in DIFFICULTIES:
        r.bad("difficulty must be easy, medium or hard")
    r.difficulty = d if d in DIFFICULTIES else None
    if raw["question_text"] and len(raw["question_text"].strip()) < 15:
        r.bad("question_text is too short (at least 15 characters)")
    if qt == "MCQ":
        opts = [raw[f"option_{c.lower()}"] for c in LETTERS]
        filled = [bool(o) for o in opts]
        if any(filled[i] is False and any(filled[i + 1:]) for i in range(4)):
            r.bad("Options must be filled in order (no gaps between option_a and option_e)")
        options = [o for o in opts if o]
        if len(options) >= 2 and len({o.lower() for o in options}) != len(options):
            r.bad("Options must be distinct")
        r.options = options
        letter = raw["correct_option"].strip().upper()
        if letter:
            if letter not in LETTERS or LETTERS.index(letter) >= len(options):
                r.bad("correct_option must be a letter (A-E) that matches a filled option")
            else:
                r.correct_index = LETTERS.index(letter)
    else:
        r.criteria = _split(raw["technical_rubric"])
        r.concepts = _split(raw["expected_concepts"])
        if not r.concepts and r.criteria:
            r.concepts = list(r.criteria)
            r.bad("expected_concepts missing: derived from the rubric criteria", "warning")
        if len(r.concepts or []) < 2:
            r.bad("Provide at least 2 rubric criteria or expected concepts")
    if raw["time_limit_seconds"]:
        if not raw["time_limit_seconds"].isdigit() or not 10 <= int(raw["time_limit_seconds"]) <= 7200:
            r.bad("time_limit_seconds should be a whole number between 10 and 7200 (ignored)", "warning")
    if raw["max_score"]:
        try:
            if not 0 < float(raw["max_score"]) <= 100:
                raise ValueError
        except ValueError:
            r.bad("max_score should be a number greater than 0 and at most 100 (ignored)", "warning")


def finalize_status(r: RowResult) -> None:
    levels = {i["level"] for i in r.issues}
    if any(i["level"] == "error" for i in r.issues):
        r.status = "INVALID"
    elif r.duplicate_of:
        r.status = "DUPLICATE"
    elif r.skill_id is None:
        r.status = "NEEDS_SKILL_MAPPING"
    elif "warning" in levels:
        r.status = "WARNING"
    else:
        r.status = "READY"


async def _org_hashes(db: AsyncSession, org_id: uuid.UUID, assessment_id: uuid.UUID | None) -> dict[str, dict]:
    """Hashes of this organization's own private bank and of the questions in its current assessment (which may include public platform ones)."""
    out: dict[str, dict] = {}
    for q in (await db.scalars(select(Question).where(Question.organization_id == org_id, Question.status.not_in([QS.REJECTED.value, QS.RETIRED.value])))).all():
        out[content_hash(q.question_text)] = {"question_id": str(q.id), "where": "your question bank"}
    if assessment_id:
        rows = (await db.scalars(select(Question).join(AssessmentQuestion, AssessmentQuestion.question_id == Question.id)
                                 .where(AssessmentQuestion.assessment_id == assessment_id))).all()
        for q in rows:
            out.setdefault(content_hash(q.question_text), {"question_id": str(q.id), "where": "this assessment"})
    return out


async def validate_rows(db: AsyncSession, org_id: uuid.UUID, raw_rows: list[dict], assessment_id: uuid.UUID | None) -> list[RowResult]:
    results = [RowResult(row=i + 2, raw=r) for i, r in enumerate(raw_rows)]  # sheet row numbers (header is row 1)
    known = await _org_hashes(db, org_id, assessment_id)
    seen_in_file: dict[str, int] = {}
    scope = TenantScope(organization_id=org_id)
    for r in results:
        check_structure(r)
        if r.raw["skill"]:
            sid, _ = await normalize_skill_name(db, r.raw["skill"])
            if sid is not None:
                skill = await db.get(Skill, sid)
                r.skill_id, r.canonical_skill = str(sid), skill.canonical_name
                if skill.canonical_name.strip().lower() != r.raw["skill"].strip().lower():
                    r.bad(f"'{r.raw['skill']}' was mapped to {skill.canonical_name}", "info")
        if r.raw["question_text"] and not any(i["level"] == "error" for i in r.issues):
            h = content_hash(r.raw["question_text"])
            if h in seen_in_file:
                r.duplicate_of = {"where": "this file", "row": seen_in_file[h]}
                r.bad(f"Same question text as row {seen_in_file[h]} in this file", "duplicate")
            elif h in known:
                r.duplicate_of = known[h]
                r.bad(f"Already exists in {known[h]['where']}", "duplicate")
            else:
                seen_in_file[h] = r.row
                if r.skill_id:  # semantic near-duplicate: this organization's scope only
                    res = validate_structure(QuestionType(r.question_type), r.raw["question_text"], r.difficulty or "medium", r.concepts,
                                             {"criteria": r.criteria} if r.criteria else None, r.options, r.correct_index, None)
                    sem = await asyncio.to_thread(validate_semantics, res, r.raw["question_text"], r.canonical_skill, uuid.UUID(r.skill_id), scope)
                    if sem.duplicate_of:
                        r.duplicate_of = {"where": "your question bank", "question_id": sem.duplicate_of}
                        r.bad("Near-duplicate of a question you already have", "duplicate")
                    elif not sem.checks.get("skill_alignment", True):
                        r.bad(f"The question text looks weakly related to {r.canonical_skill}", "warning")
        finalize_status(r)
    return results


async def revalidate_one(db: AsyncSession, org_id: uuid.UUID, batch: QuestionImportBatch, idx: int, skill_id: uuid.UUID | None) -> None:
    """Recruiter resolved a skill mapping: rebuild that row's result."""
    raws = [r["raw"] for r in batch.rows]
    saved = dict(batch.rows[idx])
    res = (await validate_rows(db, org_id, [raws[idx]], batch.assessment_id))[0]
    if skill_id is not None:
        skill = await db.get(Skill, skill_id)
        res.skill_id, res.canonical_skill = str(skill_id), skill.canonical_name
        res.issues = [i for i in res.issues if i["message"] != f"'{res.raw['skill']}' was mapped to {skill.canonical_name}"]
        res.bad(f"'{res.raw['skill']}' was mapped to {skill.canonical_name} by you", "info")
        finalize_status(res)
    res.row = saved["row"]
    res.excluded = saved.get("excluded", False)
    batch.rows = [*batch.rows[:idx], res.as_dict(), *batch.rows[idx + 1:]]


def summarize(results: list[dict]) -> dict:
    live = [r for r in results if not r.get("excluded")]
    return {"total_rows": len(results), "valid_rows": sum(r["status"] in ("READY", "WARNING") for r in live),
            "invalid_rows": sum(r["status"] in ("INVALID", "NEEDS_SKILL_MAPPING") for r in live),
            "duplicate_rows": sum(r["status"] == "DUPLICATE" for r in live)}


# ---------------------------------------------------------------- confirm
async def confirm_batch(db: AsyncSession, batch: QuestionImportBatch, user_id: uuid.UUID) -> dict:
    """Persist every non-excluded READY/WARNING row as a COMPANY_PRIVATE question of the batch's organization."""
    known = await _org_hashes(db, batch.organization_id, None)
    created, skipped_dup = [], 0
    for r in batch.rows:
        if r.get("excluded") or r["status"] not in ("READY", "WARNING"):
            continue
        raw = r["raw"]
        h = content_hash(raw["question_text"])
        if h in known:  # became a duplicate since the preview (another batch)
            skipped_dup += 1
            continue
        qtype = QuestionType(r["question_type"])
        meta = {"external_question_id": raw["external_question_id"] or None, "explanation": raw["explanation"] or None, "tags": _split(raw["tags"]) or None,
                "time_limit_seconds": int(raw["time_limit_seconds"]) if raw["time_limit_seconds"].isdigit() and 10 <= int(raw["time_limit_seconds"]) <= 7200 else None,
                "required": _truthy(raw["required"]), "max_score": None, "batch_row": r["row"]}
        try:
            ms = float(raw["max_score"]) if raw["max_score"] else None
            meta["max_score"] = ms if ms and 0 < ms <= 100 else None
        except ValueError:
            pass
        q = Question(question_text=raw["question_text"].strip(), question_type=qtype, skill_id=uuid.UUID(r["skill_id"]), difficulty=r["difficulty"],
                     options=r["options"] if qtype == QuestionType.MCQ else None, correct_option_index=r["correct_index"] if qtype == QuestionType.MCQ else None,
                     expected_concepts=r["concepts"] if qtype == QuestionType.TECHNICAL else None,
                     rubric={"criteria": r["criteria"], "version": "rubric_v1", "imported": True} if qtype == QuestionType.TECHNICAL else None,
                     source_type=QuestionSourceType.COMPANY_PRIVATE, organization_id=batch.organization_id, visibility=Visibility.COMPANY_PRIVATE,
                     status=QS.VALIDATED, content_hash=h, created_by_user_id=user_id, provenance="COMPANY_IMPORT", import_batch_id=batch.id,
                     import_meta=meta, validation_report={"ok": True, "checks": {"import_validated": True}, "reasons": [i["message"] for i in r["issues"]]})
        db.add(q)
        await db.flush()
        await asyncio.to_thread(index_question, q)  # tenant-scoped point (organization_id in the payload)
        known[h] = {"question_id": str(q.id), "where": "your question bank"}
        created.append(q)
    batch.status, batch.imported_rows = "CONFIRMED", len(created)
    batch.confirmed_at = dt.datetime.now(dt.timezone.utc)
    return {"imported": len(created), "skipped_duplicates": skipped_dup, "question_ids": [str(q.id) for q in created]}


# ---------------------------------------------------------------- coverage
async def coverage_for_job(db: AsyncSession, org_id: uuid.UUID, job, assessment) -> dict:
    """Needed (deterministic blueprint) vs available (this company's usable private questions, plus approved platform ones) vs selected
    (in the current draft assessment). Generation only has to fill what is neither available nor selected."""
    from app.models.assessments import Assessment
    from app.services.assessments.generator import blueprint_for, get_confirmed_job_skills
    from app.services.questions.governance import USABLE_OWN_COMPANY, USABLE_PLATFORM

    comps = await get_confirmed_job_skills(db, job.id)
    if not comps:
        return {"rows": [], "needed": 0, "covered": 0, "gap": 0, "coverage_percentage": None, "note": "Confirm the job requirements to see coverage."}
    bp = blueprint_for(comps, job.assessment_target_questions)
    names = {c["skill_id"]: c["skill_name"] for c in comps}
    selected: dict[tuple, int] = {}
    if assessment is not None:
        for qtype, sid in (await db.execute(select(Question.question_type, Question.skill_id).join(AssessmentQuestion, AssessmentQuestion.question_id == Question.id)
                                            .where(AssessmentQuestion.assessment_id == assessment.id))).all():
            selected[(str(sid), str(qtype))] = selected.get((str(sid), str(qtype)), 0) + 1
    rows, needed_total, covered_total, gap_total = [], 0, 0, 0
    for a in bp.allocations:
        for qt, need in (("MCQ", a.mcq_count), ("TECHNICAL", a.technical_count), ("CODING", a.coding_count)):
            if not need:
                continue
            own = len((await db.scalars(select(Question.id).where(Question.organization_id == org_id, Question.skill_id == a.skill_id, Question.question_type == qt,
                                                                   Question.status.in_([s.value for s in USABLE_OWN_COMPANY])))).all())
            plat = len((await db.scalars(select(Question.id).where(Question.visibility == Visibility.PLATFORM_PUBLIC, Question.skill_id == a.skill_id, Question.question_type == qt,
                                                                    Question.status.in_([s.value for s in USABLE_PLATFORM])))).all())
            sel = selected.get((str(a.skill_id), qt), 0)
            needed_total += need
            covered_total += min(need, max(sel, min(own + plat, need)))
            gap_total += max(0, need - max(sel, own + plat))
            rows.append({"skill": names.get(str(a.skill_id), a.skill_name), "type": qt, "needed": need, "available_private": own, "available_platform": plat, "selected": min(sel, need) if sel else 0})
    return {"rows": rows, "needed": needed_total, "covered": covered_total, "gap": gap_total,
            "coverage_percentage": round(100 * covered_total / needed_total, 1) if needed_total else None}
